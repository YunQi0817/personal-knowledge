"""PowerPoint 解析器（``.pptx``）。

幻灯片边界是**文件格式强制**的，识别不会出错——**一张幻灯片就是一个语义单元**，
比从纯文本里猜章节可靠得多。

三个有意为之的处理：

1. **正文缩进还原大纲**：``paragraph.level``（0~8）是 PowerPoint 的缩进层级，
   转成 Markdown 列表缩进，保留原有的大纲感。
2. **备注页（讲稿）拼进正文**：很多 PPT 的完整解释只写在备注里，幻灯片上
   只有几个关键词。丢掉备注等于丢掉最主要的内容。
3. **表格转 Markdown**：PowerPoint 表格常是参数对比表，转成 Markdown 后
   行列关系仍在，不会退化成空格分隔的一串词。

已知限制：整页是图片的幻灯片在这里提不出文字，需要 OCR（当前未接入）。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Union

from pptx import Presentation

from app.services.parsers.base import Chunk, SectionBuilder

#: 本解析器负责的扩展名
SUFFIXES = {".pptx"}

#: 缩进层级 -> Markdown 列表前缀
_BULLET = ("- ", "  - ", "    - ", "      - ", "        - ")


def _lines_of(text_frame: Any) -> list[str]:
    """取出文本框的段落，按缩进层级还原成 Markdown 列表行。"""
    lines: list[str] = []
    for paragraph in text_frame.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        level = min(max(int(paragraph.level or 0), 0), len(_BULLET) - 1)
        lines.append(_BULLET[level] + text)
    return lines


def _table_to_markdown(table: Any) -> str:
    """表格转 Markdown；表头取第一行。"""
    rows = [
        [(cell.text or "").strip().replace("\n", " ") for cell in row.cells]
        for row in table.rows
    ]
    rows = [row for row in rows if any(row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = "| " + " | ".join(rows[0]) + " |"
    rule = "| " + " | ".join("---" for _ in range(width)) + " |"
    return "\n".join([header, rule, *("| " + " | ".join(r) + " |" for r in rows[1:])])


def _slide_title(slide: Any) -> str:
    """取幻灯片标题：优先标题占位符，退回第一个非空文本框的首行。"""
    title_shape = slide.shapes.title
    if title_shape is not None and title_shape.has_text_frame:
        text = (title_shape.text_frame.text or "").strip()
        if text:
            return text.splitlines()[0].strip()
    for shape in slide.shapes:
        if shape.has_text_frame and (shape.text_frame.text or "").strip():
            return shape.text_frame.text.strip().splitlines()[0].strip()
    return ""


def parse(path: Union[str, Path]) -> list[Chunk]:
    """解析演示文稿，**每张幻灯片一块**（含备注页）。"""
    target = Path(path)
    presentation = Presentation(str(target))
    builder = SectionBuilder(source=target.name)

    for index, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        # 注意：必须比 shape_id，不能用 `shape is title_shape`。
        # `shapes.title` 与遍历 `shapes` 得到的是**不同的 Python 对象**，
        # 只是包装了同一个 XML 元素，`is` 判断恒为 False，
        # 结果标题会被当成正文再输出一遍。
        title_id = title_shape.shape_id if title_shape is not None else None
        title = _slide_title(slide)
        builder.heading(1, title or f"第 {index} 页", page=index)

        for shape in slide.shapes:
            if title_id is not None and shape.shape_id == title_id:
                continue
            if getattr(shape, "has_table", False):
                markdown = _table_to_markdown(shape.table)
                if markdown:
                    builder.body(markdown, content_type="table", page=index)
            elif shape.has_text_frame:
                lines = _lines_of(shape.text_frame)
                if lines:
                    builder.body("\n".join(lines), page=index)

        # 备注页：注意 has_notes_slide，直接访问 notes_slide 会凭空创建一个
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                builder.body(f"【备注】{notes}", page=index)

    return builder.chunks()
