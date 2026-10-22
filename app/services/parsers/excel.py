"""Excel 解析器（``.xlsx`` / ``.xlsm``）。

**唯一的铁律：绝不要把二维表格原样塞进向量。** 每一行都要转成
「列名：值」的自然语言：

    套餐名称：基础版；价格：99；成员人数上限：10

理由与「基础版多少钱」那次教训同源——要让向量看到**语义**，
而不是一串裸露的数字。实测把表格拍平成 ``试用版 0 1`` 之后，
列名与值的绑定关系直接断掉，那样的文本根本没法检索。

四个处理要点：

1. **合并单元格**只有左上角有值，其余是 ``None``。必须按
   ``worksheet.merged_cells`` 把值填回区域内其余格子，否则「基础版」
   只会出现在第一行，后续行看起来像是没有套餐名。
2. **公式没有缓存值**：文件从未被 Excel 打开过时，``data_only=True``
   会把公式格全读成 ``None``，此时退回读公式字符串本身。
3. **每块重复表头**，否则第 500 行那一块看不出每列是什么。
4. **首行未必是表头**（上面常有标题行、空行），需要启发式挑一行。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Union

from openpyxl import load_workbook

from app import config
from app.services.parsers.base import Chunk, SectionBuilder

#: 本解析器负责的扩展名
SUFFIXES = {".xlsx", ".xlsm"}

#: 在前若干行里寻找表头
_HEADER_SCAN = 6

#: 单个单元格最长保留字符数
_MAX_CELL = 120


# ------------------------------------------------------------------ 读表


def _has_content(rows: list[list[Any]]) -> bool:
    return any(
        cell is not None and str(cell).strip() for row in rows for cell in row
    )


def _fill_merged(worksheet: Any, rows: list[list[Any]]) -> list[list[Any]]:
    """把合并区域左上角的值填回区域内其余格子。"""
    for merged in worksheet.merged_cells.ranges:
        if merged.min_row > len(rows):
            continue
        top = rows[merged.min_row - 1]
        value = top[merged.min_col - 1] if merged.min_col <= len(top) else None
        if value is None:
            continue
        for r in range(merged.min_row, min(merged.max_row, len(rows)) + 1):
            row = rows[r - 1]
            while len(row) < merged.max_col:
                row.append(None)
            for c in range(merged.min_col, merged.max_col + 1):
                if (r, c) != (merged.min_row, merged.min_col):
                    row[c - 1] = value
    return rows


def _read_once(path: Path, *, data_only: bool) -> list[tuple[str, list[list[Any]]]]:
    # 刻意不开 read_only：只读模式的 ReadOnlyWorksheet **没有 merged_cells**，
    # 拿不到合并区域就无法把值填回去（合并单元格是本解析器必须处理的情况）。
    # 代价是大工作簿更吃内存——个人知识库场景可以接受。
    workbook = load_workbook(str(path), data_only=data_only, read_only=False)
    try:
        sheets: list[tuple[str, list[list[Any]]]] = []
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            sheets.append((worksheet.title, _fill_merged(worksheet, rows)))
        return sheets
    finally:
        workbook.close()


def _read_sheets(path: Path) -> list[tuple[str, list[list[Any]]]]:
    """读全部工作表；公式没有缓存值时自动退回读公式字符串。"""
    sheets = _read_once(path, data_only=True)
    if any(_has_content(rows) for _, rows in sheets):
        return sheets
    return _read_once(path, data_only=False)


# ------------------------------------------------------------------ 转换


def _clip(value: Any) -> str:
    text = " ".join(str(value).split())
    return text if len(text) <= _MAX_CELL else text[:_MAX_CELL] + "…"


def _find_header(rows: list[list[Any]]) -> int:
    """在前几行里挑最像表头的一行；非空单元格多且整行都是文字者胜出。"""
    best_index, best_score = 0, -1.0
    for index, row in enumerate(rows[:_HEADER_SCAN]):
        cells = [c for c in row if c is not None and str(c).strip()]
        if len(cells) < 2:
            continue
        text_ratio = sum(1 for c in cells if isinstance(c, str)) / len(cells)
        score = len(cells) * text_ratio
        if score > best_score:
            best_index, best_score = index, score
    return best_index


def _row_to_text(header: list[Any], row: list[Any]) -> str:
    """一行数据 -> 「列名：值；列名：值」。"""
    parts: list[str] = []
    for index, value in enumerate(row):
        if value is None or str(value).strip() == "":
            continue
        key = header[index] if index < len(header) else None
        key = str(key).strip() if key is not None and str(key).strip() else f"列{index + 1}"
        parts.append(f"{key}：{_clip(value)}")
    return "；".join(parts)


# ------------------------------------------------------------------ 入口


def parse(path: Union[str, Path]) -> list[Chunk]:
    """解析工作簿，**一个工作表一个章节**，行按预算聚块并重复表头。"""
    target = Path(path)
    # merge_tiny=False：Excel 的每个行组都是**刻意切开**的，
    # 若把"只有一行"的 sheet 挂起，会被并进下一张工作表里，两张表就混了。
    builder = SectionBuilder(source=target.name, merge_tiny=False)

    for title, raw_rows in _read_sheets(target):
        rows = [row for row in raw_rows if _has_content([row])]
        if not rows:
            continue

        builder.heading(1, title)
        header_index = _find_header(rows)
        header = rows[header_index]
        data = rows[header_index + 1 :]
        header_line = "；".join(
            _clip(c) for c in header if c is not None and str(c).strip()
        )

        budget = max(80, config.CHUNK_SIZE - len(header_line) - 1)
        pending: list[str] = []
        used = 0
        emitted = False

        for row in data:
            line = _row_to_text(header, row)
            if not line:
                continue
            if pending and used + len(line) > budget:
                builder.body("\n".join([header_line, *pending]), content_type="table")
                builder.flush()
                pending, used, emitted = [], 0, True
            pending.append(line)
            used += len(line) + 1

        if pending:
            builder.body("\n".join([header_line, *pending]), content_type="table")
            builder.flush()
            emitted = True

        if not emitted and header_line:
            # 只有表头没有数据行，或所有数据行都是空的
            builder.body(header_line, content_type="table")
            builder.flush()

    return builder.chunks()
